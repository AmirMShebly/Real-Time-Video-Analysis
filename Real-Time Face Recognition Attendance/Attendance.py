import os
from datetime import datetime
from openpyxl import load_workbook, Workbook
import numpy as np
import pandas as pd
import cv2
import face_recognition
from datetime import datetime


path = 'Employee_images'
images = []
class_names = []
myList = os.listdir(path)
for cls in os.listdir(path):
    img = cv2.imread(f'{path}/{cls}')
    images.append(img)
    class_names.append(os.path.splitext(cls)[0])


def encodings(images):
    encodings = []
    for img in images:
        img = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
        encoding = face_recognition.face_encodings(img)[0]
        encodings.append(encoding)
    
    return encodings


def attendance(name, filepath='attendance.xlsx', sheet_name='Sheet1'):
    """Records attendance in an xlsx file. Creates the file if it doesn't exist."""
    try:
        wb = load_workbook(filepath)
        ws = wb[sheet_name]
    except FileNotFoundError:
        wb = Workbook()
        ws = wb.active
        ws.append(['Name', 'Time']) 

    names = [row[0] for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None] 

    if name not in names:
        now = datetime.now()
        time_str = now.strftime('%H:%M:%S')
        ws.append([name, time_str])
        wb.save(filepath)
        print(f"Attendance for {name} recorded at {time_str}")
    else:
      print(f"{name} already present.")



encodings_known = encodings(images)

cap = cv2.VideoCapture(0)

while True:
    suc, ret = cap.read()
    img = cv2.resize(img, (0, 0), None, 0.25, 0.25)
    img = cv2.cvtColor(ret, cv2.COLOR_BGR2RGB)

    faces = face_recognition.face_locations(img)
    encodings = face_recognition.face_encodings(img, faces)

    for encface, face in zip(encodings, faces):
        matches = face_recognition.compare_faces(encodings_known, encface)
        face_distance = face_recognition.face_distance(encodings_known, encface)
        match_idx = np.argmin(face_distance)

        if matches[match_idx]:
            name = class_names[match_idx]
            y1, x2, y2, x1 = face
            cv2.rectangle(img, (x1, y1), (x2, y2), (0, 255, 0), 3)
            cv2.rectangle(img, (x1, y2-35), (x2, y2), (0, 255, 0), cv2.FILLED)
            cv2.putText(img, name, (x1+6, y2-6), cv2.FONT_HERSHEY_COMPLEX, 1, (255, 255, 255), 2)
            attendance(name)

    

    cv2.imshow("Image", img)
    cv2.waitKey(1)







